from __future__ import annotations

import argparse
import re
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


IMAGE_RE = re.compile(r"^!\[([^]]*)\]\((.+)\)$")
LINK_RE = re.compile(r"\[([^]]+)\]\((https?://[^)]+)\)")
INLINE_RE = re.compile(r"\[([^]]+)\]\((https?://[^)]+)\)|\*\*(.+?)\*\*")
NUMBERED_RE = re.compile(r"^(\d+)\.\s+(.+)$")


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def set_left_border(paragraph, color: str, size: str = "16") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), color)
    borders.append(left)


def set_bottom_border(paragraph, color: str = "B8C8C5") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(116, 127, 124)
    set_east_asia_font(run, "Microsoft YaHei")


def create_decimal_numbering(document: Document, start_value: int = 1) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), str(start_value))
    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.extend([start, num_format, level_text, level_justification])

    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "280")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    paragraph_properties.extend([tabs, indent, spacing])
    level.append(paragraph_properties)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "167A72")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run_properties.extend([fonts, color, underline])
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline_content(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_east_asia_font(run, "Microsoft YaHei")
        if match.group(1) is not None:
            add_hyperlink(paragraph, match.group(1), match.group(2))
        else:
            run = paragraph.add_run(match.group(3))
            run.bold = True
            set_east_asia_font(run, "Microsoft YaHei")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_east_asia_font(run, "Microsoft YaHei")


def configure_styles(document: Document) -> None:
    # Design preset: narrative_proposal, with a named CJK font override.
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(43, 50, 48)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title_properties = title.element.get_or_add_pPr()
    border = title_properties.find(qn("w:pBdr"))
    if border is not None:
        title_properties.remove(border)
    empty_border = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right", "between", "bar"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "nil")
        empty_border.append(edge)
    title_properties.append(empty_border)

    for style_name, size, color in (
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.keep_with_next = True

    styles["Heading 1"].paragraph_format.space_before = Pt(18)
    styles["Heading 1"].paragraph_format.space_after = Pt(10)
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)
    styles["Heading 3"].paragraph_format.space_before = Pt(8)
    styles["Heading 3"].paragraph_format.space_after = Pt(4)

    if "Article Quote" not in styles:
        quote = styles.add_style("Article Quote", WD_STYLE_TYPE.PARAGRAPH)
    else:
        quote = styles["Article Quote"]
    quote.font.name = "Microsoft YaHei"
    quote._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    quote.font.size = Pt(10)
    quote.font.color.rgb = RGBColor(72, 83, 80)
    quote.paragraph_format.left_indent = Cm(0.5)
    quote.paragraph_format.right_indent = Cm(0.3)
    quote.paragraph_format.space_before = Pt(4)
    quote.paragraph_format.line_spacing = 1.208
    quote.paragraph_format.space_after = Pt(8)


def add_image(document: Document, image_path: Path, alt_text: str) -> None:
    if not image_path.exists():
        raise FileNotFoundError(image_path)

    with Image.open(image_path) as image:
        width_px, height_px = image.size
    max_width = 15.2
    max_height = 10.5
    scale = min(max_width / width_px, max_height / height_px)
    width_cm = width_px * scale
    height_cm = height_px * scale

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm), height=Cm(height_cm))


def add_body_paragraph(
    document: Document,
    text: str,
    numbered: bool = False,
    num_id: int | None = None,
):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = None if numbered else Cm(0.74)
    if numbered:
        num_properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number_id = OxmlElement("w:numId")
        number_id.set(qn("w:val"), str(num_id))
        num_properties.extend([level, number_id])
    add_inline_content(paragraph, text)
    return paragraph


def build_document(markdown_path: Path, readable_closing: bool = True, balanced_title: bool = True,
                   output_root: Path | None = None, author: str | None = None) -> Path:
    from _article_tables import read_table
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    configure_styles(document)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    header_run = header.add_run("微信公众号文章｜商丘金蓬")
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor(116, 127, 124)
    set_east_asia_font(header_run, "Microsoft YaHei")

    title_text = ""
    index = 0
    active_num_id: int | None = None
    in_closing = False
    while index < len(lines):
        raw = lines[index].strip()
        index += 1
        if not raw:
            continue

        parsed_table = read_table(lines, index - 1)
        if parsed_table:
            rows, index = parsed_table
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = False
            for row_index, values in enumerate(rows):
                row = table.add_row()
                properties = row._tr.get_or_add_trPr()
                properties.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    properties.append(OxmlElement("w:tblHeader"))
                for cell, value in zip(row.cells, values):
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.first_line_indent = None
                    paragraph.paragraph_format.line_spacing = 1.1
                    paragraph.paragraph_format.space_before = Pt(4)
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.keep_with_next = row_index == 0
                    add_inline_content(paragraph, value)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "0B1D33" if row_index == 0 else ("F3F6F8" if row_index % 2 else "FFFFFF"))
                    cell._tc.get_or_add_tcPr().append(shading)
            document.add_paragraph().paragraph_format.space_after = Pt(0)
            active_num_id = None
            continue

        if raw.startswith("# "):
            active_num_id = None
            title_text = raw[2:].strip()
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if balanced_title and "：" in title_text:
                lead, detail = title_text.split("：", 1)
                add_inline_content(paragraph, lead + "：")
                paragraph.add_run().add_break()
                add_inline_content(paragraph, detail)
            else:
                add_inline_content(paragraph, title_text)
            continue

        if raw.startswith("## "):
            active_num_id = None
            heading_text = raw[3:].strip()
            in_closing = heading_text == "关于商丘金蓬"
            paragraph = document.add_paragraph(heading_text, style="Heading 1")
            continue

        if raw.startswith("#### "):
            active_num_id = None
            document.add_paragraph(raw[5:].strip(), style="Heading 3")
            continue

        if raw.startswith("### "):
            active_num_id = None
            document.add_paragraph(raw[4:].strip(), style="Heading 2")
            continue

        if raw == "---":
            active_num_id = None
            separator = document.add_paragraph()
            set_bottom_border(separator)
            continue

        image_match = IMAGE_RE.match(raw)
        if image_match:
            active_num_id = None
            relative = Path(unquote(image_match.group(2)))
            add_image(document, markdown_path.parent / relative, image_match.group(1))
            continue

        if raw.startswith("> "):
            active_num_id = None
            quote_lines = [raw[2:].strip()]
            while index < len(lines) and lines[index].strip().startswith("> "):
                quote_lines.append(lines[index].strip()[2:].strip())
                index += 1
            paragraph = document.add_paragraph(style="Article Quote")
            add_inline_content(paragraph, " ".join(quote_lines))
            set_cell_shading(paragraph, "F1F6F5")
            set_left_border(paragraph, "3F8F85")
            continue

        numbered_match = NUMBERED_RE.match(raw)
        if numbered_match:
            active_num_id = create_decimal_numbering(document, int(numbered_match.group(1)))
            add_body_paragraph(
                document,
                numbered_match.group(2),
                numbered=True,
                num_id=active_num_id,
            )
            continue

        active_num_id = None
        paragraph_lines = [raw]
        while index < len(lines):
            next_line = lines[index].strip()
            if not next_line:
                index += 1
                break
            if (
                next_line.startswith("#")
                or next_line.startswith("> ")
                or next_line == "---"
                or IMAGE_RE.match(next_line)
                or NUMBERED_RE.match(next_line)
            ):
                break
            paragraph_lines.append(next_line)
            index += 1
        text = " ".join(paragraph_lines)
        if text.startswith("- "):
            text = text[2:]
        paragraph = add_body_paragraph(document, text)
        if in_closing:
            paragraph.paragraph_format.line_spacing = 0.9
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)

    if readable_closing:
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if re.match(r"^表\d+\s", text) or (len(text) < 60 and text.endswith("：")):
                paragraph.paragraph_format.keep_with_next = True
        closing_start = next((i for i, p in enumerate(document.paragraphs) if p.text == "关于商丘金蓬"), None)
        if closing_start is not None:
            closing_paragraphs = document.paragraphs[closing_start + 1:]
            for i, paragraph in enumerate(closing_paragraphs):
                paragraph.paragraph_format.line_spacing = 1.2
                paragraph.paragraph_format.space_after = Pt(5)
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.keep_with_next = i < len(closing_paragraphs) - 1
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)

    document.core_properties.title = title_text
    document.core_properties.subject = "商丘金蓬微信公众号文章"
    author_match = re.search(r"（修改后）([^（）]+)$", markdown_path.stem)
    document.core_properties.author = author if author is not None else (author_match.group(1) if author_match else "")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("商丘金蓬｜第")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(130, 142, 139)
    set_east_asia_font(footer_run, "Microsoft YaHei")
    add_page_field(footer)
    page_suffix = footer.add_run("页")
    page_suffix.font.size = Pt(8)
    page_suffix.font.color.rgb = RGBColor(130, 142, 139)
    set_east_asia_font(page_suffix, "Microsoft YaHei")

    output_root = output_root or markdown_path.resolve().parent.parent / "修改后Word"
    output_root.mkdir(parents=True, exist_ok=True)
    output_path = output_root / f"{markdown_path.stem}.docx"
    document.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Build reviewed Word articles from Markdown.")
    parser.add_argument("markdown", nargs="+", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    markdown_files = [path.resolve() for path in args.markdown]
    for markdown_path in markdown_files:
        output = build_document(markdown_path, output_root=args.output_dir)
        print(output.name)


if __name__ == "__main__":
    main()
