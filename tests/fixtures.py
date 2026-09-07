"""Synthetic test material, never a production article or user image."""
import json
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from article_tool import CLOSING, build, extract


def create_fixture(root: Path):
    root.mkdir(parents=True, exist_ok=True)
    picture = root / "synthetic.png"
    Image.new("RGB", (800, 240), (116, 200, 172)).save(picture)
    title = "物料与工艺：一篇用于检查排版的合成文章"
    source = root / "原稿.docx"
    document = Document()
    document.add_heading(title, 0)
    document.add_paragraph("这份合成样例包含 4,010 个测试记录，用来检查数字、文字和表格能否完整保留。")
    document.add_heading("项目产能与产品亮点", 1)
    for item in range(1, 7):
        document.add_paragraph(f"检查项{item}：第{item}项正文应当保持可见。")
    document.add_picture(str(picture), width=Inches(5))
    document.add_paragraph("图1 合成测试图")
    document.add_heading("表格与原料", 1)
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["对象", "指标", "备注"]):
        cell.text = value
    for row in [["测试原料", "40%", "用于验证百分号完整显示"], ["测试设备", "1.84", "小数不应被截开"]]:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value
    document.add_paragraph("文字与原图均应完整保留。")
    document.save(source)
    snapshot = extract(source, root / "source-extract")
    data = json.loads(snapshot.read_text(encoding="utf-8"))
    assets = root / "Markdown" / "assets"
    assets.mkdir(parents=True)
    copied = assets / "synthetic.png"
    shutil.copy2(data["images"][0]["path"], copied)
    md = root / "Markdown" / f"2026.09.07{title}（修改后）李明.md"
    content = [f"# {title}", "这份合成样例包含 **4,010** 个测试记录，用来检查数字、文字和表格能否完整保留。",
               "## 项目产能与产品亮点"]
    content += [f"- **检查项{item}：** 第{item}项正文应当保持可见。" for item in range(1, 7)]
    content += ["![测试图片](assets/synthetic.png)", "图1 合成测试图", "## 表格与原料",
                "| 对象 | 指标 | 备注 |\n| --- | --- | --- |\n| 测试原料 | 40% | 用于验证百分号完整显示 |\n| 测试设备 | 1.84 | 小数不应被截开 |",
                "文字与原图均应完整保留。", CLOSING]
    md.write_text("\n\n".join(content) + "\n", encoding="utf-8")
    audit = root / "审查记录" / "fixture.md"
    audit.parent.mkdir()
    audit.write_text("内部合成测试，非真实业务文章。用于测试机械校验，不声称完成专业审查。", encoding="utf-8")
    manifest = build(md, snapshot, audit, root)
    return manifest
